from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Literal

import sqlglot
from openpyxl import load_workbook
from sqlglot import exp


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
SCHEMA_PATH = DATA_DIR / "schema_demo.json"
STATISTICS_PATH = DATA_DIR / "statistics_demo.xlsx"
SCHEDULE_PATH = DATA_DIR / "horarios_ejecucion.json"
POLICIES_PATH = DATA_DIR / "politicas_tecnicas_queries.json"

ALLOWED_PROFILES = {"analista", "operador", "dba"}
ALLOWED_INTENTIONS = {"consulta", "descarga", "proceso_batch"}
SAFE_IDENTIFIER = re.compile(r"^[A-Za-z0-9_.@-]{3,128}$")
SAFE_SCHEMA = re.compile(r"^[A-Za-z_][A-Za-z0-9_]{0,127}$")
DENIED_NODES = (
    exp.Insert, exp.Update, exp.Delete, exp.Merge, exp.Create, exp.Drop,
    exp.Alter, exp.Command, exp.Transaction, exp.Commit, exp.Rollback,
    exp.Grant, exp.Revoke,
)


def _error(code: str, message: str) -> dict[str, Any]:
    return {
        "ok": False,
        "error": {"code": code, "message": message},
        "decision": "rejected",
        "requires_dba_approval": False,
        "evidence_complete": False,
        "missing_evidence": [],
        "warnings": [],
    }


def _load_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(path.name)
    return json.loads(path.read_text(encoding="utf-8"))


def _load_statistics(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "tables" not in workbook.sheetnames:
        raise ValueError("statistics_demo.xlsx no contiene la hoja tables")
    rows = list(workbook["tables"].iter_rows(values_only=True))
    required = {"schema_name", "table_name", "approx_row_count", "statistics_updated_at"}
    header_index = next(
        (index for index, row in enumerate(rows[:20]) if required.issubset({str(value).strip() for value in row if value is not None})),
        None,
    )
    if header_index is None:
        raise ValueError("No se encontró la cabecera requerida en statistics_demo.xlsx")
    headers = [str(value).strip() if value is not None else "" for value in rows[header_index]]
    records: list[dict[str, Any]] = []
    for row in rows[header_index + 1 :]:
        if not row or all(value is None for value in row):
            continue
        record = dict(zip(headers, row, strict=False))
        for key, value in tuple(record.items()):
            if isinstance(value, datetime):
                record[key] = value.isoformat()
        records.append(record)
    return records


def _minutes(value: str) -> int:
    hours, minutes = map(int, value.split(":"))
    return hours * 60 + minutes


def _schedule_allows(query_class: str, execution_time: str, schedule: dict[str, Any]) -> tuple[bool, dict[str, Any] | None]:
    window = next((item for item in schedule.get("execution_windows", []) if item.get("query_class") == query_class), None)
    if not window:
        return False, None
    current, start, end = _minutes(execution_time), _minutes(window["start_time"]), _minutes(window["end_time"])
    allowed = (current >= start or current < end) if window.get("crosses_midnight") else start <= current < end
    return allowed, window


def analyze_risk_query(
    query: str,
    usuario: str,
    perfil_usuario: Literal["analista", "operador", "dba"],
    hora_ejecucion: str,
    intencion: Literal["consulta", "descarga", "proceso_batch"] = "consulta",
    schema_name: str = "query_analyzer",
) -> dict[str, Any]:
    """Analiza una consulta de lectura sin ejecutarla y devuelve evidencia estructurada."""
    query = query.strip() if isinstance(query, str) else ""
    usuario = usuario.strip() if isinstance(usuario, str) else ""
    schema_name = schema_name.strip() if isinstance(schema_name, str) else ""
    if not query:
        return _error("EMPTY_QUERY", "La consulta SQL es obligatoria.")
    if len(query) > 20_000:
        return _error("QUERY_TOO_LONG", "La consulta supera el límite de 20000 caracteres.")
    if not SAFE_IDENTIFIER.fullmatch(usuario):
        return _error("INVALID_USER", "El usuario tiene un formato inválido.")
    if perfil_usuario not in ALLOWED_PROFILES:
        return _error("INVALID_PROFILE", "El perfil no está autorizado.")
    if intencion not in ALLOWED_INTENTIONS:
        return _error("INVALID_INTENTION", "La intención no está autorizada.")
    if not SAFE_SCHEMA.fullmatch(schema_name):
        return _error("INVALID_SCHEMA", "El esquema tiene un formato inválido.")
    if not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", hora_ejecucion or ""):
        return _error("INVALID_TIME", "La hora debe usar el formato HH:MM.")

    try:
        statements = sqlglot.parse(query, read="postgres")
    except (sqlglot.errors.ParseError, ValueError):
        return _error("INVALID_SQL", "La consulta no puede analizarse como PostgreSQL.")
    if len(statements) != 1:
        return _error("MULTIPLE_STATEMENTS", "Se permite exactamente una sentencia.")
    root = statements[0]
    denied = next((node for node in root.walk() if isinstance(node, DENIED_NODES)), None)
    if denied is not None or (not isinstance(root, (exp.Select, exp.Union, exp.Intersect, exp.Except)) and root.find(exp.Select) is None):
        return _error("WRITE_OPERATION_DENIED", "Solo se analizan consultas de lectura; no se ejecutó ninguna operación.")

    ctes = {cte.alias_or_name.lower() for cte in root.find_all(exp.CTE)}
    tables: list[dict[str, str | None]] = []
    seen: set[tuple[str, str]] = set()
    for table in root.find_all(exp.Table):
        if table.name.lower() in ctes:
            continue
        schema = table.db or schema_name
        if schema.lower() != schema_name.lower():
            return _error("UNAUTHORIZED_SCHEMA", f"El esquema {schema} no está autorizado.")
        key = (schema.lower(), table.name.lower())
        if key not in seen:
            seen.add(key)
            tables.append({"schema_name": schema, "table_name": table.name, "alias": table.alias or None})

    has_where = any(select.args.get("where") is not None for select in root.find_all(exp.Select))
    join_count = sum(len(select.args.get("joins") or []) for select in root.find_all(exp.Select))
    uses_temp = bool(root.find(exp.Into)) or bool(root.find(exp.TemporaryProperty))
    has_aggregation = any(isinstance(node, exp.AggFunc) for node in root.walk())
    query_type = "select_aggregate" if has_aggregation or root.find(exp.Group) else "select_detail"
    heavy_reasons: list[str] = []
    if len(tables) > 1:
        heavy_reasons.append("multiple_tables")
    if join_count:
        heavy_reasons.append("join")
    if root.find(exp.Group):
        heavy_reasons.append("group_by")
    if root.find(exp.Order):
        heavy_reasons.append("order_by")
    if root.find(exp.Distinct):
        heavy_reasons.append("distinct")
    if not has_where:
        heavy_reasons.append("without_where")
    if uses_temp:
        heavy_reasons.append("temporary_table")
    query_class = "pesado" if heavy_reasons else "liviano"

    try:
        schema_doc = _load_json(SCHEMA_PATH)
        policies = _load_json(POLICIES_PATH)
        schedule = _load_json(SCHEDULE_PATH)
        statistics = _load_statistics(STATISTICS_PATH)
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        result = _error("SOURCE_UNAVAILABLE", "No fue posible leer una fuente de evidencia autorizada.")
        result["warnings"] = [type(exc).__name__]
        return result

    metadata: list[dict[str, Any]] = []
    selected_statistics: list[dict[str, Any]] = []
    missing: list[str] = []
    for table in tables:
        schema, name = table["schema_name"].lower(), table["table_name"].lower()
        metadata_item = next((item for item in schema_doc.get("objects", []) if str(item.get("schema", "")).lower() == schema and str(item.get("name", "")).lower() == name), None)
        statistic_item = next((item for item in statistics if str(item.get("schema_name", "")).lower() == schema and str(item.get("table_name", "")).lower() == name), None)
        if metadata_item:
            metadata.append(metadata_item)
        else:
            missing.append(f"schema_metadata:{schema}.{name}")
        if statistic_item:
            selected_statistics.append(statistic_item)
        else:
            missing.append(f"statistics:{schema}.{name}")

    profile = policies.get("profiles", {}).get(perfil_usuario, {})
    registered_user = next((item for item in policies.get("users", []) if item.get("user_id") == usuario), None)
    incidents = [item for item in policies.get("query_incidents", []) if item.get("user_id") == usuario]
    security_incidents = [item for item in policies.get("security_incidents", []) if item.get("user_id") == usuario]
    alerts: list[dict[str, str]] = []

    def alert(policy_id: str, severity: str, reason: str) -> None:
        alerts.append({"policy_id": policy_id, "severity": severity, "reason": reason})

    if not registered_user or registered_user.get("status") != "active" or registered_user.get("profile") != perfil_usuario:
        alert("POL-QRY-001", "critical", "Usuario inexistente, inactivo o con perfil inconsistente.")
    if query_type not in profile.get("allowed_query_types", []):
        alert("POL-QRY-001", "critical", "Tipo de consulta no permitido para el perfil.")
    max_table_rows = max((int(item.get("approx_row_count") or 0) for item in selected_statistics), default=0)
    if not has_where and max_table_rows > int(profile.get("max_estimated_rows_scanned", 0)):
        alert("POL-VOL-001", "high", f"Tabla sin filtro con {max_table_rows} filas aproximadas; supera el límite del perfil.")
    if intencion == "descarga" and not has_where and max_table_rows > int(profile.get("max_download_rows", 0)):
        alert("POL-DWN-001", "critical", "Descarga sin filtro por encima del límite del perfil.")
    if uses_temp:
        alert("POL-TMP-001", "high", "La consulta intenta crear una tabla temporal.")
    if len(tables) > int(profile.get("max_joined_tables", 0)):
        alert("POL-JOIN-001", "high", "El cruce excede el máximo de tablas del perfil.")
    if not has_where and max_table_rows > 50_000:
        alert("POL-SCAN-001", "high", "Posible full scan sobre una tabla voluminosa.")
    if len(incidents) >= 2:
        alert("POL-INC-001", "high", "El usuario registra incidentes operativos recurrentes.")
    if any(item.get("status") == "open" or item.get("severity") == "critical" for item in security_incidents):
        alert("POL-SEC-001", "critical", "El usuario registra un incidente de seguridad abierto o crítico.")

    schedule_allowed, window = _schedule_allows(query_class, hora_ejecucion, schedule)
    requires_approval = bool(alerts) or not schedule_allowed
    evidence_complete = not missing
    decision = "requires_dba_approval" if requires_approval else ("approved" if evidence_complete else "incomplete")
    masked = root.transform(lambda node: exp.Placeholder() if isinstance(node, exp.Literal) else node).sql(dialect="postgres")
    return {
        "ok": True,
        "decision": decision,
        "requires_dba_approval": requires_approval,
        "evidence_complete": evidence_complete,
        "missing_evidence": missing,
        "warnings": ["La clasificación es estática y no reemplaza EXPLAIN."],
        "source": {
            "schema_metadata": SCHEMA_PATH.name,
            "statistics": STATISTICS_PATH.name,
            "schedule": SCHEDULE_PATH.name,
            "technical_policies": POLICIES_PATH.name,
        },
        "request_context": {"usuario": usuario, "perfil_usuario": perfil_usuario, "hora_ejecucion": hora_ejecucion, "intencion": intencion},
        "query_analysis": {
            "masked_query": masked,
            "query_type": query_type,
            "query_class": query_class,
            "heavy_reasons": heavy_reasons,
            "tables": tables,
            "has_where": has_where,
            "join_count": join_count,
            "uses_temporary_table": uses_temp,
        },
        "evidence": {
            "metadata": metadata,
            "statistics": selected_statistics,
            "profile_limits": profile,
            "operational_incident_count": len(incidents),
            "open_or_critical_security_incident": any(item.get("status") == "open" or item.get("severity") == "critical" for item in security_incidents),
        },
        "policy_alerts": alerts,
        "schedule": {"allowed": schedule_allowed, "applied_window": window, "timezone": schedule.get("timezone")},
    }

